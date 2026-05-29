from docx import Document
from docx.shared import RGBColor

from engine.formatter.docx_formatter import apply_profile_formatting
from engine.models import ParagraphStyle, Profile, PageSetup
from engine.profiles.loader import load_profile
from engine.structure.models import RecognizedParagraph, RecognizedStructure, RecognizedTable


def test_apply_profile_formatting_sets_title_and_body(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.docx"

    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("这是正文段落。")
    doc.save(source)

    apply_profile_formatting(source, output, load_profile("general"))

    formatted = Document(output)
    assert formatted.paragraphs[0].alignment == 1
    assert formatted.paragraphs[0].runs[0].font.name == "宋体"
    assert formatted.paragraphs[1].runs[0].font.name == "仿宋_GB2312"


def test_heading_3_detection_and_formatting(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("一、一级标题")
    doc.add_paragraph("（一）二级标题")
    doc.add_paragraph("1. 三级标题")
    doc.add_paragraph("正文内容")
    doc.save(source)

    profile = Profile(
        id="test",
        name="test",
        title=ParagraphStyle(font="宋体", size="二号", bold=True, align="center"),
        body=ParagraphStyle(font="仿宋_GB2312", size="三号", align="justify", first_line_indent="2em", line_spacing="29pt"),
        headings={
            "heading_1": ParagraphStyle(font="黑体", size="三号", bold=True, line_spacing="29pt"),
            "heading_2": ParagraphStyle(font="楷体_GB2312", size="三号", bold=True, line_spacing="29pt"),
            "heading_3": ParagraphStyle(font="仿宋_GB2312", size="三号", bold=True, line_spacing="29pt"),
        },
    )

    apply_profile_formatting(source, output, profile)

    result = Document(output)
    h3 = result.paragraphs[3]
    assert h3.runs[0].font.name == "仿宋_GB2312"
    assert h3.runs[0].font.bold is True


def test_date_and_department_field_formatting(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("文稿标题")
    doc.add_paragraph("2026年5月27日")
    doc.add_paragraph("某某部门")
    doc.add_paragraph("正文内容")
    doc.save(source)

    profile = Profile(
        id="test",
        name="test",
        title=ParagraphStyle(font="宋体", size="二号", bold=True, align="center"),
        body=ParagraphStyle(font="仿宋_GB2312", size="三号", align="justify"),
        date_field=ParagraphStyle(font="楷体_GB2312", size="三号", bold=True, align="center"),
        department_field=ParagraphStyle(font="楷体_GB2312", size="三号", bold=True, align="center"),
    )

    apply_profile_formatting(source, output, profile)

    result = Document(output)
    assert result.paragraphs[1].runs[0].font.name == "楷体_GB2312"
    assert result.paragraphs[1].runs[0].font.bold is True
    assert result.paragraphs[1].alignment == 1  # CENTER
    assert result.paragraphs[2].runs[0].font.name == "楷体_GB2312"
    assert result.paragraphs[2].runs[0].font.bold is True
    assert result.paragraphs[2].alignment == 1  # CENTER


def test_apply_profile_formatting_uses_recognized_paragraph_roles(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("文稿标题")
    doc.add_paragraph("这不是第二段日期")
    doc.add_paragraph("2026年5月29日")
    doc.add_paragraph("某某部门")
    doc.add_paragraph("一、总体要求")
    doc.add_paragraph("正文内容")
    doc.save(source)

    structure = RecognizedStructure(
        paragraphs=[
            RecognizedParagraph(index=0, role="title", confidence=0.9, reason="title"),
            RecognizedParagraph(index=1, role="body", confidence=0.8, reason="body"),
            RecognizedParagraph(index=2, role="date", confidence=0.9, reason="date"),
            RecognizedParagraph(index=3, role="department", confidence=0.9, reason="department"),
            RecognizedParagraph(index=4, role="heading_1", confidence=0.9, reason="heading"),
            RecognizedParagraph(index=5, role="body", confidence=0.8, reason="body"),
        ],
        tables=[],
    )

    apply_profile_formatting(source, output, load_profile("default"), structure=structure)

    result = Document(output)
    assert result.paragraphs[2].runs[0].font.name == "楷体_GB2312"
    assert result.paragraphs[2].alignment == 1
    assert result.paragraphs[3].runs[0].font.name == "楷体_GB2312"
    assert result.paragraphs[4].runs[0].font.name == "黑体"
    assert result.paragraphs[5].runs[0].font.name == "仿宋_GB2312"


def test_apply_profile_formatting_styles_table_cell_text(tmp_path):
    source = tmp_path / "table.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "序号"
    table.cell(0, 1).text = "事项"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "起草方案"
    doc.save(source)

    structure = RecognizedStructure(
        paragraphs=[RecognizedParagraph(index=0, role="title", confidence=0.9, reason="title")],
        tables=[RecognizedTable(index=0, role="data_table", confidence=0.9, reason="table")],
    )

    apply_profile_formatting(source, output, load_profile("default"), structure=structure)

    result = Document(output)
    run = result.tables[0].cell(1, 1).paragraphs[0].runs[0]
    assert run.font.name == "仿宋_GB2312"
    assert run.font.color.rgb == RGBColor(0, 0, 0)

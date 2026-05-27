from docx import Document

from engine.diagnostics.diagnose import diagnose_docx
from engine.fixer.apply_fixes import apply_fixes
from engine.profiles.loader import load_profile


def test_diagnose_default_profile(tmp_path):
    docx_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("日期")
    doc.add_paragraph("部门")
    doc.add_paragraph("一、一级标题")
    doc.add_paragraph("（一）二级标题")
    doc.add_paragraph("1. 三级标题")
    doc.add_paragraph("正文")
    doc.save(docx_path)

    profile = load_profile("default")
    issues = diagnose_docx(docx_path, profile)

    # Should detect font mismatches since we used default fonts
    assert len(issues) > 0
    # Check that heading_3 is checked
    h3_issues = [i for i in issues if "3" in i.message and "段" in i.message]
    # There should be issues about paragraphs not matching expected fonts
    assert any("仿宋_GB2312" in i.message or "楷体_GB2312" in i.message or "黑体" in i.message for i in issues)


def test_diagnose_and_fix_body_font(tmp_path):
    source = tmp_path / "bad.docx"
    fixed = tmp_path / "fixed.docx"

    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文段落")
    doc.save(source)

    profile = load_profile("general")
    issues = diagnose_docx(source, profile)

    assert any(issue.fix_action for issue in issues)

    apply_fixes(source, fixed, issues)
    remaining = diagnose_docx(fixed, profile)

    assert len([issue for issue in remaining if issue.severity == "error"]) == 0

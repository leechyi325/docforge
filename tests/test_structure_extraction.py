from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from engine.structure.extract import extract_structure_input


def test_extract_structure_input_reads_paragraphs_and_numbering_hints(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    title = doc.add_paragraph("关于开展测试工作的通知")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("补充").bold = True
    doc.add_paragraph("2026年5月29日")
    doc.add_paragraph("一、总体要求")
    doc.add_paragraph("（一）工作目标")
    doc.add_paragraph("1. 具体安排")
    doc.save(source)

    structure_input = extract_structure_input(source)

    assert structure_input.source_path == str(source)
    assert structure_input.paragraphs[0].index == 0
    assert structure_input.paragraphs[0].text == "关于开展测试工作的通知补充"
    assert structure_input.paragraphs[0].alignment == "center"
    assert structure_input.paragraphs[0].is_bold is True
    assert structure_input.paragraphs[2].numbering_hint == "chinese_heading_1"
    assert structure_input.paragraphs[3].numbering_hint == "chinese_heading_2"
    assert structure_input.paragraphs[4].numbering_hint == "numbered_heading_3"


def test_extract_structure_input_reads_table_shape_and_samples(tmp_path):
    source = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("表1 工作安排")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "序号"
    table.cell(0, 1).text = "事项"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "起草方案"
    table.cell(2, 0).text = "2"
    table.cell(2, 1).text = "汇总意见"
    doc.add_paragraph("说明：以上为初步安排")
    doc.save(source)

    structure_input = extract_structure_input(source)

    assert len(structure_input.tables) == 1
    candidate = structure_input.tables[0]
    assert candidate.index == 0
    assert candidate.rows == 3
    assert candidate.columns == 2
    assert candidate.before_paragraph_index == 1
    assert candidate.after_paragraph_index == 2
    assert candidate.caption_paragraph_index == 1
    assert candidate.sample_cells[0] == ["序号", "事项"]
    assert candidate.sample_cells[1] == ["1", "起草方案"]

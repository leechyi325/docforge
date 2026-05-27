from __future__ import annotations

from pathlib import Path
from typing import List

from docx import Document

from engine.models import FixAction, Issue, Location, ParagraphStyle, Profile


def diagnose_docx(input_docx: Path, profile: Profile) -> List[Issue]:
    document = Document(input_docx)
    issues: List[Issue] = []

    if not document.paragraphs:
        return [
            Issue(
                id="issue_empty_document",
                severity="error",
                message="文档没有段落内容",
                location=Location(type="document", index=None, key=None),
                fix_action=None,
            )
        ]

    for index, paragraph in enumerate(document.paragraphs):
        expected = _get_expected_style(index, paragraph.text, profile)
        if expected is None:
            continue

        run = paragraph.runs[0] if paragraph.runs else None
        actual_font = run.font.name if run else None

        if expected.font and actual_font != expected.font:
            issues.append(
                Issue(
                    id=f"issue_paragraph_{index}_font",
                    severity="warning",
                    message=f"第 {index + 1} 段字体不符合要求，应为 {expected.font}",
                    location=Location(type="paragraph", index=index, key=None),
                    fix_action=FixAction(
                        type="set_paragraph_format",
                        target=Location(type="paragraph", index=index, key=None),
                        properties=ParagraphStyle(font=expected.font, size=expected.size, align=expected.align),
                    ),
                )
            )

    return issues


def _get_expected_style(index: int, text: str, profile: Profile):
    if index == 0:
        return profile.title
    if profile.date_field and index == 1:
        return profile.date_field
    if profile.department_field and index == 2:
        return profile.department_field

    from engine.formatter.docx_formatter import _detect_heading_key
    heading_key = _detect_heading_key(text)
    if heading_key and heading_key in profile.headings:
        return profile.headings[heading_key]

    return profile.body

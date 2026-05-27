from docx import Document
from docx.shared import Pt

from engine.formatter.docx_formatter import apply_profile_formatting
from engine.models import PageNumbering, PageSetup, ParagraphStyle, Profile


def _make_profile_with_page_number():
    return Profile(
        id="test",
        name="test",
        title=ParagraphStyle(font="宋体", size="二号", bold=True, align="center"),
        body=ParagraphStyle(font="仿宋", size="三号"),
        page=PageSetup(
            margin_top="3.5cm",
            margin_bottom="3.5cm",
            margin_left="2.8cm",
            margin_right="2.8cm",
            page_number=PageNumbering(
                font="宋体",
                size="四号",
                format="－ {page} －",
                odd_align="right",
                even_align="left",
            ),
        ),
    )


def test_page_numbering_creates_footer_paragraphs(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(source)

    profile = _make_profile_with_page_number()
    apply_profile_formatting(source, output, profile)

    result = Document(output)
    section = result.sections[0]
    footer = section.footer
    assert len(footer.paragraphs) >= 1


def test_page_numbering_font_and_size(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(source)

    profile = _make_profile_with_page_number()
    apply_profile_formatting(source, output, profile)

    result = Document(output)
    footer = result.sections[0].footer
    para = None
    for p in footer.paragraphs:
        if p.runs:
            para = p
            break
    assert para is not None
    assert para.runs[0].font.name == "宋体"
    assert para.runs[0].font.size == Pt(14)  # 四号 = 14pt

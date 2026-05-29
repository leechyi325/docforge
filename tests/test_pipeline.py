import shutil

import pytest
from docx import Document

import engine.pipeline
from engine.pipeline import UnsupportedInputError, format_document
from engine.profiles.loader import load_profile


def test_format_document_accepts_docx_input(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文段落")
    doc.save(source)

    result = format_document(source, output, load_profile("general"))

    assert result.output_path == output
    assert output.exists()
    formatted = Document(output)
    assert formatted.paragraphs[0].runs[0].font.name == "宋体"
    assert formatted.paragraphs[1].runs[0].font.name == "仿宋_GB2312"
    assert isinstance(result.issues, list)


def test_format_document_rejects_doc_input(tmp_path):
    source = tmp_path / "legacy.doc"
    output = tmp_path / "formatted.docx"
    source.write_bytes(b"not a real doc file")

    with pytest.raises(UnsupportedInputError, match="Markdown and .docx"):
        format_document(source, output, load_profile("general"))


def test_format_document_requires_docx_output(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.pdf"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.save(source)

    with pytest.raises(ValueError, match="Output path must end with .docx"):
        format_document(source, output, load_profile("general"))


def test_format_document_reports_missing_input(tmp_path):
    source = tmp_path / "missing.docx"
    output = tmp_path / "formatted.docx"

    with pytest.raises(FileNotFoundError, match="Input file does not exist"):
        format_document(source, output, load_profile("general"))


def test_format_document_requires_new_output_path(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.save(source)

    with pytest.raises(ValueError, match="Output path must be different"):
        format_document(source, source, load_profile("general"))


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is required for markdown conversion")
def test_format_document_accepts_markdown_input(tmp_path):
    source = tmp_path / "source.md"
    output = tmp_path / "formatted.docx"
    source.write_text("# 测试标题\n\n正文段落\n", encoding="utf-8")

    result = format_document(source, output, load_profile("general"))

    assert result.output_path == output
    assert output.exists()
    assert Document(output).paragraphs[0].text.strip() == "测试标题"


import json
import sys


def test_format_document_emits_progress_events(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文")
    doc.save(source)

    captured = []

    original_stderr = sys.stderr

    class FakeStderr:
        def write(self, data):
            stripped = data.strip()
            if stripped.startswith("{"):
                try:
                    captured.append(json.loads(stripped))
                except json.JSONDecodeError:
                    pass
            return len(data)
        def flush(self):
            pass

    monkeypatch.setattr(sys, "stderr", FakeStderr())

    format_document(source, output, load_profile("general"))

    stages = [e.get("stage") for e in captured if "stage" in e]
    assert "formatting" in stages
    assert "diagnosing" in stages


from engine.structure.models import RecognizedParagraph, RecognizedStructure, RecognizedTable


class FakeStructureClient:
    def __init__(self):
        self.inputs = []

    def recognize_structure(self, structure_input):
        self.inputs.append(structure_input)
        return RecognizedStructure(
            paragraphs=[
                RecognizedParagraph(index=0, role="title", confidence=0.9, reason="first paragraph"),
                RecognizedParagraph(index=1, role="date", confidence=0.9, reason="date"),
                RecognizedParagraph(index=2, role="department", confidence=0.9, reason="department"),
                RecognizedParagraph(index=3, role="body", confidence=0.8, reason="body"),
            ],
            tables=[],
        )


class FailingStructureClient:
    def recognize_structure(self, structure_input):
        raise ValueError("AI 智能识别需要配置远程模型")


def test_format_document_default_profile_requires_structure_client(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.save(source)

    with pytest.raises(ValueError, match="AI 智能识别需要配置远程模型"):
        format_document(source, output, load_profile("default"))


def test_format_document_default_profile_requires_structure_client_before_markdown_conversion(monkeypatch, tmp_path):
    source = tmp_path / "source.md"
    output = tmp_path / "output.docx"
    source.write_text("# 标题\n\n正文\n", encoding="utf-8")

    def fail_if_called(input_path, output_path):
        raise AssertionError("pandoc should not be called")

    monkeypatch.setattr(engine.pipeline, "convert_markdown_to_docx", fail_if_called)

    with pytest.raises(ValueError, match="AI 智能识别需要配置远程模型"):
        format_document(source, output, load_profile("default"))


def test_format_document_default_profile_uses_structure_recognition(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("文稿标题")
    doc.add_paragraph("2026年5月29日")
    doc.add_paragraph("某某部门")
    doc.add_paragraph("正文内容")
    doc.save(source)
    fake_client = FakeStructureClient()

    result = format_document(source, output, load_profile("default"), llm_client=fake_client)

    assert output.exists()
    assert len(fake_client.inputs) == 1
    assert result.structure_summary is not None
    assert result.structure_summary.counts["title"] == 1
    assert result.structure_summary.counts["date"] == 1
    assert result.structure_summary.counts["department"] == 1


def test_format_document_general_profile_does_not_call_structure_recognition(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(source)

    result = format_document(source, output, load_profile("general"), llm_client=FailingStructureClient())

    assert output.exists()
    assert result.structure_summary is None


def test_format_document_rejects_invalid_recognized_structure_index(tmp_path):
    class BadIndexClient:
        def recognize_structure(self, structure_input):
            return RecognizedStructure(
                paragraphs=[RecognizedParagraph(index=99, role="title", confidence=0.9, reason="bad")],
                tables=[],
            )

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.save(source)

    with pytest.raises(ValueError, match="paragraph index 99"):
        format_document(source, output, load_profile("default"), llm_client=BadIndexClient())


def test_format_document_default_profile_summarizes_tables(tmp_path):
    class TableClient:
        def recognize_structure(self, structure_input):
            return RecognizedStructure(
                paragraphs=[RecognizedParagraph(index=0, role="title", confidence=0.9, reason="title")],
                tables=[RecognizedTable(index=0, role="data_table", confidence=0.9, reason="table")],
            )

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "内容"
    doc.save(source)

    result = format_document(source, output, load_profile("default"), llm_client=TableClient())

    assert result.structure_summary is not None
    assert result.structure_summary.counts["table"] == 1

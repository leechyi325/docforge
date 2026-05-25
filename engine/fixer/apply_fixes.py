from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Set

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from engine.models import Issue, ParagraphStyle
from engine.style_utils import chinese_size_to_pt

ALLOWED_ACTIONS = {"set_paragraph_format", "apply_heading_style", "set_section_format"}
ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def apply_fixes(
    input_docx: Path,
    output_docx: Path,
    issues: List[Issue],
    selected_issue_ids: Optional[Set[str]] = None,
) -> Path:
    document = Document(input_docx)
    selected = selected_issue_ids or {issue.id for issue in issues}

    for issue in issues:
        if issue.id not in selected or issue.fix_action is None:
            continue
        action = issue.fix_action
        if action.type not in ALLOWED_ACTIONS:
            continue
        if action.target.type != "paragraph" or action.target.index is None:
            continue
        if action.target.index >= len(document.paragraphs):
            continue
        if not isinstance(action.properties, ParagraphStyle):
            continue
        _apply_paragraph_properties(document.paragraphs[action.target.index], action.properties)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    return output_docx


def _apply_paragraph_properties(paragraph, style: ParagraphStyle) -> None:
    if style.align:
        paragraph.alignment = ALIGNMENT[style.align]
    runs = paragraph.runs or [paragraph.add_run("")]
    for run in runs:
        if style.font:
            run.font.name = style.font
            run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", style.font)
        if style.size:
            run.font.size = Pt(chinese_size_to_pt(style.size))

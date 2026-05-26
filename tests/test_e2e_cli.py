import shutil

import pytest
from docx import Document

from engine.cli import main


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is required for e2e conversion")
def test_generate_general_docx_from_markdown(tmp_path):
    input_md = tmp_path / "input.md"
    output_docx = tmp_path / "output.docx"
    input_md.write_text(
        "# 测试标题\n\n## 一、总体要求\n\n这是正文段落。\n",
        encoding="utf-8",
    )

    exit_code = main(
        [
            "generate",
            "--input",
            str(input_md),
            "--profile",
            "general",
            "--format-instruction",
            "标题用方正小标宋二号居中，正文仿宋三号，行距固定28磅",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 0
    assert output_docx.exists()
    doc = Document(output_docx)
    assert doc.paragraphs[0].text.strip() == "测试标题"


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is required for e2e conversion")
def test_format_general_docx_from_markdown(tmp_path):
    input_md = tmp_path / "input.md"
    output_docx = tmp_path / "output.docx"
    input_md.write_text(
        "# 测试标题\n\n## 一、总体要求\n\n这是正文段落。\n",
        encoding="utf-8",
    )

    exit_code = main(
        [
            "format",
            "--input",
            str(input_md),
            "--profile",
            "general",
            "--format-instruction",
            "标题用方正小标宋二号居中，正文仿宋三号，行距固定28磅",
            "--llm-provider",
            "local",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 0
    assert output_docx.exists()
    doc = Document(output_docx)
    assert doc.paragraphs[0].text.strip() == "测试标题"

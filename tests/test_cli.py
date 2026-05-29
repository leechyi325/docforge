import json

from docx import Document

import engine.cli
from engine.cli import main


def test_cli_diagnose_outputs_json(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    exit_code = main(["diagnose", "--input", str(input_docx), "--profile", "general"])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert "issues" in payload


def test_cli_format_docx_outputs_json_and_writes_file(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    exit_code = main(["format", "--input", str(input_docx), "--profile", "general", "--output", str(output_docx)])

    assert exit_code == 0
    assert output_docx.exists()
    payload = json.loads(capsys.readouterr().out)
    assert payload["output"] == str(output_docx)
    assert "issues" in payload


def test_cli_format_accepts_openai_responses_provider_without_call_when_no_instruction(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    exit_code = main(
        [
            "format",
            "--input",
            str(input_docx),
            "--profile",
            "general",
            "--llm-provider",
            "openai-responses",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 0
    assert output_docx.exists()
    payload = json.loads(capsys.readouterr().out)
    assert payload["output"] == str(output_docx)


def test_cli_format_accepts_openai_alias_without_call_when_no_instruction(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    exit_code = main(
        [
            "format",
            "--input",
            str(input_docx),
            "--profile",
            "general",
            "--llm-provider",
            "openai",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 0
    assert output_docx.exists()
    payload = json.loads(capsys.readouterr().out)
    assert payload["output"] == str(output_docx)


def test_cli_format_remote_provider_requires_model_when_instruction_is_present(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    exit_code = main(
        [
            "format",
            "--input",
            str(input_docx),
            "--profile",
            "general",
            "--format-instruction",
            "正文仿宋三号",
            "--llm-provider",
            "openai-compatible",
            "--api-key",
            "test",
            "--llm-base-url",
            "https://api.example.test/v1",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "requires --llm-model" in captured.err


def test_cli_format_reports_missing_input(tmp_path, capsys):
    output_docx = tmp_path / "output.docx"

    exit_code = main(["format", "--input", str(tmp_path / "missing.docx"), "--output", str(output_docx)])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Input file does not exist" in captured.err


def test_cli_format_reports_model_runtime_error(monkeypatch, tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("测试标题")
    doc.save(input_docx)

    class FailingClient:
        def parse_format_instruction(self, instruction):
            raise RuntimeError("provider failed")

    monkeypatch.setattr(engine.cli, "create_llm_client", lambda settings: FailingClient())

    exit_code = main(
        [
            "format",
            "--input",
            str(input_docx),
            "--profile",
            "general",
            "--format-instruction",
            "正文仿宋三号",
            "--llm-provider",
            "openai",
            "--api-key",
            "test",
            "--output",
            str(output_docx),
        ]
    )

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "provider failed" in captured.err


def test_cli_format_rejects_doc_input(tmp_path, capsys):
    input_doc = tmp_path / "legacy.doc"
    output_docx = tmp_path / "output.docx"
    input_doc.write_bytes(b"not a real doc file")

    exit_code = main(["format", "--input", str(input_doc), "--profile", "general", "--output", str(output_docx)])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Markdown and .docx" in captured.err


from engine.structure.models import RecognizedParagraph, RecognizedStructure


def test_cli_format_default_with_local_provider_fails(tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.save(input_docx)

    exit_code = main([
        "format",
        "--input",
        str(input_docx),
        "--profile",
        "default",
        "--llm-provider",
        "local",
        "--output",
        str(output_docx),
    ])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "AI 智能识别需要配置远程模型" in captured.err


def test_cli_format_default_outputs_structure_summary(monkeypatch, tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)

    class FakeClient:
        def parse_format_instruction(self, instruction):
            raise AssertionError("format instruction parser should not be called")

        def recognize_structure(self, structure_input):
            return RecognizedStructure(
                paragraphs=[
                    RecognizedParagraph(index=0, role="title", confidence=0.9, reason="title"),
                    RecognizedParagraph(index=1, role="body", confidence=0.8, reason="body"),
                ],
                tables=[],
            )

    monkeypatch.setattr(engine.cli, "create_llm_client", lambda settings: FakeClient())

    exit_code = main([
        "format",
        "--input",
        str(input_docx),
        "--profile",
        "default",
        "--llm-provider",
        "openai-compatible",
        "--llm-model",
        "fake-model",
        "--api-key",
        "fake-key",
        "--output",
        str(output_docx),
    ])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["structure_summary"]["title"] == 1
    assert payload["structure_summary"]["body"] == 1


def test_cli_format_default_with_instruction_reuses_one_client(monkeypatch, tmp_path, capsys):
    input_docx = tmp_path / "input.docx"
    output_docx = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("标题")
    doc.add_paragraph("正文")
    doc.save(input_docx)
    calls = []
    events = []

    class FakeClient:
        def parse_format_instruction(self, instruction):
            events.append(("parse", instruction))
            return None

        def recognize_structure(self, structure_input):
            events.append(("recognize", len(structure_input.paragraphs)))
            return RecognizedStructure(
                paragraphs=[
                    RecognizedParagraph(index=0, role="title", confidence=0.9, reason="title"),
                    RecognizedParagraph(index=1, role="body", confidence=0.8, reason="body"),
                ],
                tables=[],
            )

    def fake_create_llm_client(settings):
        calls.append(settings)
        return FakeClient()

    monkeypatch.setattr(engine.cli, "create_llm_client", fake_create_llm_client)

    exit_code = main([
        "format",
        "--input",
        str(input_docx),
        "--profile",
        "default",
        "--format-instruction",
        "正文仿宋三号",
        "--llm-provider",
        "openai-compatible",
        "--llm-model",
        "fake-model",
        "--api-key",
        "fake-key",
        "--output",
        str(output_docx),
    ])

    assert exit_code == 0
    assert len(calls) == 1
    assert events == [("parse", "正文仿宋三号"), ("recognize", 2)]
    payload = json.loads(capsys.readouterr().out)
    assert payload["structure_summary"]["title"] == 1
    assert payload["structure_summary"]["body"] == 1

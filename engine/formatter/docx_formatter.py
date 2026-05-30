from __future__ import annotations

from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from engine.models import FormatOverride, PageSetup, ParagraphStyle, Profile
from engine.structure.models import RecognizedStructure
from engine.style_utils import chinese_size_to_pt, merge_style, parse_length_to_pt

ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def apply_profile_formatting(
    input_docx: Path,
    output_docx: Path,
    profile: Profile,
    override: Optional[FormatOverride] = None,
    structure: Optional[RecognizedStructure] = None,
) -> Path:
    document = Document(input_docx)
    page = override.page if override and override.page else profile.page
    _apply_page_setup(document, page)
    if page.page_number:
        _apply_page_numbering(document, page.page_number)

    title_style = merge_style(profile.title, override.title if override else None)
    body_style = merge_style(profile.body, override.body if override else None)
    heading_overrides = override.headings if override else {}

    if structure:
        role_by_index = {p.index: p.role for p in structure.paragraphs}
        for index, paragraph in enumerate(document.paragraphs):
            role = role_by_index.get(index, "body")
            style = _get_style_for_role(role, profile, title_style, body_style, heading_overrides)
            _apply_paragraph_style(paragraph, style)
            if role in ("body", "unknown") and profile.content_bold:
                _apply_content_bold(paragraph, profile.content_bold)
    else:
        for index, paragraph in enumerate(document.paragraphs):
            if index == 0:
                _apply_paragraph_style(paragraph, title_style)
            elif profile.date_field and index == 1:
                _apply_paragraph_style(paragraph, profile.date_field)
            elif profile.department_field and index == 2:
                _apply_paragraph_style(paragraph, profile.department_field)
            else:
                heading_key = _detect_heading_key(paragraph.text)
                if heading_key:
                    base = profile.headings.get(heading_key, body_style)
                    _apply_paragraph_style(paragraph, merge_style(base, heading_overrides.get(heading_key)))
                else:
                    _apply_paragraph_style(paragraph, body_style)
                    if profile.content_bold:
                        _apply_content_bold(paragraph, profile.content_bold)

    _apply_table_formatting(document, body_style)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    return output_docx


def _apply_page_setup(document: Document, page: PageSetup) -> None:
    for section in document.sections:
        section.top_margin = Cm(parse_length_to_pt(page.margin_top) / 28.3464567)
        section.bottom_margin = Cm(parse_length_to_pt(page.margin_bottom) / 28.3464567)
        section.left_margin = Cm(parse_length_to_pt(page.margin_left) / 28.3464567)
        section.right_margin = Cm(parse_length_to_pt(page.margin_right) / 28.3464567)


def _get_style_for_role(
    role: str,
    profile: Profile,
    title_style: ParagraphStyle,
    body_style: ParagraphStyle,
    heading_overrides: dict,
) -> ParagraphStyle:
    if role == "title":
        return title_style
    if role == "date" and profile.date_field:
        return profile.date_field
    if role == "department" and profile.department_field:
        return profile.department_field
    if role in ("heading_1", "heading_2", "heading_3"):
        base = profile.headings.get(role, body_style)
        return merge_style(base, heading_overrides.get(role))
    if role in profile.special_sections:
        return profile.special_sections[role]
    return body_style


def _apply_table_formatting(document: Document, body_style: ParagraphStyle) -> None:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _apply_paragraph_style(paragraph, body_style)


def _apply_paragraph_style(paragraph, style: ParagraphStyle) -> None:
    if style.align:
        paragraph.alignment = ALIGNMENT[style.align]
    if style.first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(parse_length_to_pt(style.first_line_indent))
    if style.line_spacing:
        paragraph.paragraph_format.line_spacing = Pt(parse_length_to_pt(style.line_spacing))
    if style.space_before:
        paragraph.paragraph_format.space_before = Pt(parse_length_to_pt(style.space_before))
    if style.space_after:
        paragraph.paragraph_format.space_after = Pt(parse_length_to_pt(style.space_after))

    runs = paragraph.runs or [paragraph.add_run("")]
    for run in runs:
        if style.font:
            run.font.name = style.font
            run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", style.font)
        if style.size:
            run.font.size = Pt(chinese_size_to_pt(style.size))
        if style.bold is not None:
            run.font.bold = style.bold
        run.font.italic = False
        run.font.underline = False
        run.font.color.rgb = RGBColor(0, 0, 0)


def _detect_heading_key(text: str) -> str | None:
    stripped = text.strip()
    if stripped.startswith(("一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、")):
        return "heading_1"
    if stripped.startswith(("（一）", "（二）", "（三）", "（四）", "（五）")):
        return "heading_2"
    if _looks_like_heading_3(stripped):
        return "heading_3"
    return None


def _looks_like_heading_3(text: str) -> bool:
    import re
    return bool(re.match(r"^\d+[\.\、]", text))


def _apply_page_numbering(document, page_number) -> None:
    from docx.oxml import OxmlElement

    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False

        for para in footer.paragraphs:
            for run in para.runs:
                run.clear()

        if footer.paragraphs:
            para = footer.paragraphs[0]
        else:
            para = footer.add_paragraph()

        align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        para.alignment = align_map.get(page_number.odd_align, WD_ALIGN_PARAGRAPH.RIGHT)

        run_prefix = para.add_run("－ ")
        _apply_run_font(run_prefix, page_number.font, page_number.size)

        fld_char_begin = OxmlElement("w:fldChar")
        fld_char_begin.set(qn("w:fldCharType"), "begin")

        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = " PAGE "

        fld_char_separate = OxmlElement("w:fldChar")
        fld_char_separate.set(qn("w:fldCharType"), "separate")

        fld_char_end = OxmlElement("w:fldChar")
        fld_char_end.set(qn("w:fldCharType"), "end")

        run_field = para.add_run()
        _apply_run_font(run_field, page_number.font, page_number.size)
        run_field._element.append(fld_char_begin)

        run_instr = para.add_run()
        _apply_run_font(run_instr, page_number.font, page_number.size)
        run_instr._element.append(instr_text)

        run_sep = para.add_run()
        _apply_run_font(run_sep, page_number.font, page_number.size)
        run_sep._element.append(fld_char_separate)

        run_end = para.add_run()
        _apply_run_font(run_end, page_number.font, page_number.size)
        run_end._element.append(fld_char_end)

        run_suffix = para.add_run(" －")
        _apply_run_font(run_suffix, page_number.font, page_number.size)


def _apply_run_font(run, font_name, size) -> None:
    if font_name:
        run.font.name = font_name
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name
        )
    if size:
        run.font.size = Pt(chinese_size_to_pt(size))
    run.font.color.rgb = RGBColor(0, 0, 0)


def _apply_content_bold(paragraph, bold_style: ParagraphStyle) -> None:
    for run in paragraph.runs:
        if run.font.bold:
            if bold_style.font:
                run.font.name = bold_style.font
                run._element.rPr.rFonts.set(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", bold_style.font
                )
            if bold_style.size:
                run.font.size = Pt(chinese_size_to_pt(bold_style.size))
            run.font.color.rgb = RGBColor(0, 0, 0)
